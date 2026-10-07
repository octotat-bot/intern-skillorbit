"""Fixture resumes, generated in code so tests are reproducible.

Run ``python -m tests.fixtures.resumes <out_dir>`` from backend/ to write the
files to disk for manual testing in the UI.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import docx
import pymupdf

GOOD_RESUME = """Aarav Sharma
aarav.sharma@example.com | +91 98765 43210 | linkedin.com/in/aarav-sharma | github.com/aaravsharma
Bengaluru, India

PROFESSIONAL SUMMARY
Final-year Computer Science student focused on data analysis and machine learning, with two internships and three shipped projects.

EDUCATION
B.Tech in Computer Science and Engineering
Indian Institute of Technology, Delhi | 2021 - 2025 | CGPA: 8.7/10

TECHNICAL SKILLS
Languages: Python, SQL, JavaScript, Java
Frameworks: Pandas, NumPy, scikit-learn, React, Flask
Tools: Git, Docker, Tableau, Power BI, AWS
Databases: PostgreSQL, MongoDB

EXPERIENCE
Data Analyst Intern, Flipkart | May 2024 - Jul 2024
• Automated weekly sales reporting with Python and SQL, reducing manual effort by 60%.
• Built a Tableau dashboard tracking 12 KPIs used by 3 regional teams.
Software Engineering Intern, Infosys | Jun 2023 - Aug 2023
• Developed REST APIs in Flask serving 10,000+ daily requests.

PROJECTS
Customer Churn Prediction | Python, scikit-learn, Pandas
• Engineered 25 features and trained a gradient boosting model reaching 89% accuracy.
• Deployed the model as a Flask API on AWS EC2.
Sales Analytics Dashboard | SQL, Power BI
• Designed a star-schema warehouse and dashboard analyzing 2M+ transactions.
Portfolio Website | React, Tailwind CSS
• Built a responsive portfolio with a 95+ Lighthouse performance score.

CERTIFICATIONS
• Google Data Analytics Professional Certificate (2024)
• AWS Certified Cloud Practitioner (2023)

ACHIEVEMENTS
• Winner, Smart India Hackathon 2023 (out of 500+ teams)"""

WEAK_RESUME = """RESUME

Name: Rohit Kumar
Email: rohit123@gmail.com

Objective
I am a hard worker and team player who is passionate about computers. I want a job in a good company where I can learn.

Education
BSc Computer Science

Skills
Python, MS Office

Projects
I made a website for my college fest. I helped my friends with the code."""

TWO_COLUMN_HEADER = [
    "Priya Nair",
    "priya.nair@example.com | +1 (555) 123-4567 | github.com/priyanair",
]
TWO_COLUMN_LEFT = [
    "SKILLS", "Python, SQL, Excel", "Tableau, Power BI", "Pandas, NumPy",
    "Statistics, A/B Testing", "", "EDUCATION", "B.Sc. Statistics",
    "University of Mumbai", "2020 - 2023", "", "CERTIFICATIONS",
    "Google Data Analytics", "Tableau Desktop Specialist",
]
TWO_COLUMN_RIGHT = [
    "EXPERIENCE", "Data Analyst Intern, Zomato", "Analyzed 1M+ orders to find churn drivers.",
    "Cut report turnaround time by 40%.", "Built KPI dashboards in Tableau.", "",
    "PROJECTS", "Retail Demand Forecasting", "Forecast weekly demand with Python.",
    "Reduced stockouts by 18% in a pilot.", "Survey Insights Dashboard",
    "Visualized survey data in Power BI.", "Presented to 4 stakeholders.",
]

_PAGE_MARGIN = 50
_LINE_HEIGHT_FACTOR = 1.5


def build_docx(text: str, header: str | None = None) -> bytes:
    """One paragraph per line, with optional page-header text."""
    document = docx.Document()
    if header:
        document.sections[0].header.paragraphs[0].text = header
    for line in text.split("\n"):
        document.add_paragraph(line)
    return _save_docx(document)


def build_table_docx(left: list[str], right: list[str], heading: str) -> bytes:
    """A two-cell table layout, common in resume templates."""
    document = docx.Document()
    document.add_paragraph(heading)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "\n".join(left)
    table.cell(0, 1).text = "\n".join(right)
    return _save_docx(document)


def _save_docx(document: docx.document.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(text: str, font_size: float = 10, pages: int = 1) -> bytes:
    """Single-column PDF, one text line per row, repeated across ``pages`` pages."""
    document = pymupdf.open()
    for _ in range(pages):
        page = document.new_page()
        y = _PAGE_MARGIN
        for line in text.split("\n"):
            if y > page.rect.height - _PAGE_MARGIN:
                page = document.new_page()
                y = _PAGE_MARGIN
            if line:
                page.insert_text((_PAGE_MARGIN, y), line, fontsize=font_size)
            y += font_size * _LINE_HEIGHT_FACTOR
    return document.tobytes()


def build_two_column_pdf(
    header: list[str], left: list[str], right: list[str], font_size: float = 9
) -> bytes:
    """Full-width header followed by two side-by-side columns."""
    document = pymupdf.open()
    page = document.new_page()
    step = font_size * _LINE_HEIGHT_FACTOR
    y = _PAGE_MARGIN
    for line in header:
        page.insert_text((_PAGE_MARGIN, y), line, fontsize=font_size + 1)
        y += step
    y += step
    for index in range(max(len(left), len(right))):
        for x, column in ((_PAGE_MARGIN, left), (260, right)):
            if index < len(column) and column[index]:
                page.insert_text((x, y), column[index], fontsize=font_size)
        y += step
    return document.tobytes()


def build_scanned_pdf() -> bytes:
    """A page holding only an image: what a scanned resume looks like to a parser."""
    document = pymupdf.open()
    page = document.new_page()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 300, 400), False)
    pixmap.clear_with(220)
    page.insert_image(page.rect, pixmap=pixmap)
    return document.tobytes()


def build_encrypted_pdf() -> bytes:
    document = pymupdf.open()
    document.new_page().insert_text((_PAGE_MARGIN, _PAGE_MARGIN), "Locked resume")
    return document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user"
    )


CORRUPTED_DOCX = b"PK\x03\x04" + b"this is not really a zip archive" * 10
CORRUPTED_PDF = b"%PDF-1.4\n" + b"\x00garbage bytes that are not a pdf" * 10


def write_samples(out_dir: Path) -> list[Path]:
    """Write the main fixtures to ``out_dir`` for manual upload testing."""
    out_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "good_resume.pdf": build_pdf(GOOD_RESUME),
        "good_resume.docx": build_docx(GOOD_RESUME),
        "weak_resume.docx": build_docx(WEAK_RESUME),
        "two_column_resume.pdf": build_two_column_pdf(TWO_COLUMN_HEADER, TWO_COLUMN_LEFT, TWO_COLUMN_RIGHT),
        "scanned_resume.pdf": build_scanned_pdf(),
    }
    paths = []
    for name, data in files.items():
        path = out_dir / name
        path.write_bytes(data)
        paths.append(path)
    return paths


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "sample_resumes")
    for written in write_samples(target):
        print(written)
